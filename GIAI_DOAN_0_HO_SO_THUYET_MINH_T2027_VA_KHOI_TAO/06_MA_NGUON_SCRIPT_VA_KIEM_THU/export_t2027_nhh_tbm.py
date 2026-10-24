# -*- coding: utf-8 -*-
r"""
Export the official research proposal document named: T2027_NHH_TBM.docx
According to the user's explicit request:
'xuất lại file báo cáo thuyết minh đề tài cấp trường đặt tên T2027_NHH_TBM'
"""
import os
import sys
import copy
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

sys.stdout.reconfigure(encoding='utf-8')

SRC_PATH = r"D:\OneDrive\NCKH\NCKH 2027\T2027\Insar\Thuyết minh đề tài undefined.docx"
OUT_PATHS = [
    r"D:\OneDrive\NCKH\NCKH 2027\T2027\T2027_NHH_TBM.docx",
    r"D:\OneDrive\NCKH\NCKH 2027\T2027\Insar\T2027_NHH_TBM.docx",
    r"D:\OneDrive\NCKH\NCKH 2027\T2027\TBM_ĐTM\T2027_NHH_TBM.docx"
]

def set_cell_font(cell, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color=None):
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = font_name
            run.font.size = Pt(size_pt)
            run.bold = bold
            run.italic = italic
            if color:
                run.font.color.rgb = color

def add_styled_para(cell, text, font_name="Times New Roman", size_pt=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, line_spacing=1.15):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    return p

def export_proposal():
    print(f"Loading base document from: {SRC_PATH}")
    doc = docx.Document(SRC_PATH)
    t0 = doc.tables[0]
    t1 = doc.tables[1]
    t2 = doc.tables[2]

    # 1. Update Table 0 - Code: T2027_NHH_TBM
    c_code = t0.rows[1].cells[9]
    c_code.text = ""
    add_styled_para(c_code, "2. MÃ SỐ (do Khoa cấp):\nT2027_NHH_TBM", size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 2. Update Table 0 - Acknowledgments format with grant number T2027_NHH_TBM
    c_ack = t0.rows[12].cells[7]
    c_ack.text = ""
    add_styled_para(c_ack, 
        "- Chủ nhiệm đề tài là tác giả đầu tiên hoặc tác giả liên hệ (theo email @utc.edu.vn và cơ quan làm việc chính là Trường ĐH GTVT).\n"
        "- Trong bài báo phải ghi nhận sự tài trợ trong mục “Lời cảm ơn” (“Acknowledgements”), theo mẫu:\n"
        "+ Bài báo viết bằng tiếng Anh: \"This research is funded by University of Transport and Communications (UTC) under grant number T2027_NHH_TBM\".\n"
        "+ Bài báo viết bằng tiếng Việt: “Nghiên cứu này được tài trợ bởi Trường Đại học Giao thông vận tải (ĐH GTVT) trong đề tài mã số T2027_NHH_TBM”.\n"
        "+ Đối với bài báo viết bằng ngôn ngữ khác: sử dụng cụm từ tương đương.",
        size_pt=11, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    # 3. Ensure Table 1 has all 5 tasks
    # Let's inspect schedule in t1
    tasks = [
        ("1", "Nội dung 1: Thu thập, chuẩn hóa cơ sở dữ liệu viễn thám radar Sentinel-1 (2019–2026), bình đồ, trắc dọc và số liệu địa chất, mốc quan trắc Metro 1.", "Bộ CSDL viễn thám & địa chất số hóa chuẩn GIS.", "01/2027 – 03/2027", "ThS. Nguyễn Hải Hà\nKS. Trần Văn Long"),
        ("2", "Nội dung 2: Xử lý chuỗi thời gian InSAR đa quỹ đạo (SNAP-SNAPHU-MintPy) và phân rã trường chuyển vị 2 chiều (thẳng đứng dV và ngang dEW).", "Trường vận tốc lún LOS và bản đồ phân rã chuyển vị 2D (dV, dEW).", "04/2027 – 06/2027", "ThS. Nguyễn Hải Hà\nThS. Lê Văn Nam"),
        ("3", "Nội dung 3: Đối chuẩn kết quả lún InSAR với mốc quan trắc trắc địa thực địa tại 8 mặt cắt TBM; phân tích tương quan cơ học đất và công trình ngầm.", "Báo cáo phân tích đối chuẩn thực địa, chỉ số RMSE, MAE, R².", "07/2027 – 09/2027", "ThS. Nguyễn Hải Hà\nTS. Nguyễn Thạch Bích"),
        ("4", "Nội dung 4: Phân tích phễu lún Peck (1969), đánh giá nguy cơ tổn hại công trình lân cận theo tiêu chuẩn Burland (1995) và thành lập bản đồ số GIS.", "Báo cáo đánh giá rủi ro Burland và bộ bản đồ số chuyên đề phân vùng cảnh báo (Shapefile).", "10/2027 – 11/2027", "ThS. Nguyễn Hải Hà\nThS. Lê Văn Nam"),
        ("5", "Nội dung 5: Viết bài báo khoa học đăng trên Tạp chí Giao thông vận tải (HĐGSNN); tổng hợp báo cáo tổng kết và nghiệm thu đề tài.", "01 Bài báo Tạp chí Giao thông vận tải & Báo cáo tổng kết đề tài.", "11/2027 – 12/2027", "Tập thể nhóm NCKH")
    ]

    # Find schedule header row
    header_idx = -1
    for r_i, row in enumerate(t1.rows):
        if "STT" in row.cells[0].text and "nội dung" in row.cells[1].text.lower():
            header_idx = r_i
            break

    print(f"Schedule header index in Table 1: {header_idx}")
    if header_idx != -1:
        # Check existing task rows
        current_task_rows = []
        next_sec_idx = -1
        for r_i in range(header_idx + 1, len(t1.rows)):
            first_txt = t1.rows[r_i].cells[0].text.strip()
            if first_txt in ["1", "2", "3", "4", "5"]:
                current_task_rows.append(r_i)
            else:
                next_sec_idx = r_i
                break
        
        print(f"Existing task rows count: {len(current_task_rows)}")
        # If we have 4 tasks, add Task 5
        if len(current_task_rows) == 4:
            row_task4 = t1.rows[current_task_rows[-1]]
            new_tr = copy.deepcopy(row_task4._tr)
            tcs = new_tr.xpath('.//w:tc')
            for col_idx in range(5):
                for p in tcs[col_idx].xpath('.//w:p'):
                    tcs[col_idx].remove(p)
                new_p = OxmlElement('w:p')
                new_r = OxmlElement('w:r')
                new_t = OxmlElement('w:t')
                new_t.text = tasks[4][col_idx]
                new_r.append(new_t)
                new_p.append(new_r)
                tcs[col_idx].append(new_p)
                
                # Apply styling
                rPr = OxmlElement('w:rPr')
                rFonts = OxmlElement('w:rFonts')
                rFonts.set(qn('w:ascii'), 'Times New Roman')
                rFonts.set(qn('w:hAnsi'), 'Times New Roman')
                rFonts.set(qn('w:cs'), 'Times New Roman')
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), '22')  # 11 pt
                rPr.append(rFonts)
                rPr.append(sz)
                new_r.append(rPr)
                if col_idx in (0, 3):
                    pPr = OxmlElement('w:pPr')
                    jc = OxmlElement('w:jc')
                    jc.set(qn('w:val'), 'center')
                    pPr.append(jc)
                    new_p.append(pPr)
            row_task4._tr.addnext(new_tr)
            print("Successfully appended Task 5 row to Table 1!")

    # 4. Save to all target paths
    for p in OUT_PATHS:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
        print(f"-> Successfully saved: {p} ({os.path.getsize(p)} bytes)")

if __name__ == "__main__":
    export_proposal()
