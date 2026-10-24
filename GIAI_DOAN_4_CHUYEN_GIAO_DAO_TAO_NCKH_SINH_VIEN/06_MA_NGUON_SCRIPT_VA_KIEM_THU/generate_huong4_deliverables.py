# -*- coding: utf-8 -*-
"""
Script tạo các file văn bản chính (.docx) cho đề tài NCKHSV Hướng 4:
Xây dựng bản đồ số tương tác (WebGIS/StoryMap) hỗ trợ cung cấp thông tin và tham vấn cộng đồng về vùng ảnh hưởng thi công hầm Metro Hà Nội.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TARGET_DIR = r"d:\OneDrive\NCKH\NCKH 2027\T2027\TBM_ĐTM\NCKH_SV_2026_Huong4_WebGIS_StoryMap"
os.makedirs(TARGET_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CCCCCC')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, size_pt=13, bold=True, color_rgb=(0, 51, 102))
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, size_pt=12, bold=True, color_rgb=(180, 0, 0))
    return p

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, size_pt=11.5, bold=True)
    r = p.add_run(text)
    format_run(r, size_pt=11.5, italic=italic)
    return p

# =========================================================================
# 1. TẠO 01_DE_CUONG_THUYET_MINH_NCKHSV.docx
# =========================================================================
def build_proposal_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.79)

    # Header
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_hdr.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO / TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI\n")
    format_run(r1, size_pt=11, bold=True)
    r2 = p_hdr.add_run("KHOA CÔNG TRÌNH VÀ MÔI TRƯỜNG – BỘ MÔN ĐỊA KỸ THUẬT & KHÔNG GIAN NGẦM\n")
    format_run(r2, size_pt=11, bold=True, color_rgb=(0, 51, 102))
    r3 = p_hdr.add_run("-------------------------------------\n\n")
    format_run(r3, size_pt=10)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("THUYẾT MINH ĐỀ TÀI NGHIÊN CỨU KHOA HỌC CỦA SINH VIÊN\nNĂM HỌC 2026 – 2027\n")
    format_run(r_t, size_pt=15, bold=True, color_rgb=(0, 51, 153))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("TÊN ĐỀ TÀI:\n\"XÂY DỰNG BẢN ĐỒ SỐ TƯƠNG TÁC (WEBGIS/STORYMAP) HỖ TRỢ CUNG CẤP THÔNG TIN\nVÀ THAM VẤN CỘNG ĐỒNG VỀ VÙNG ẢNH HƯỞNG THI CÔNG HẦM METRO HÀ NỘI\"\n")
    format_run(r_sub, size_pt=13, bold=True, color_rgb=(180, 0, 0))

    # General info table
    tbl = doc.add_table(rows=6, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    info_data = [
        ("Mã số đề tài:", "SV-2026-GIS-04"),
        ("Lĩnh vực nghiên cứu:", "Hệ thống Thông tin Địa lý (GIS) / Chuyển đổi số Quản lý Đô thị / Truyền thông Môi trường"),
        ("Thời gian thực hiện:", "08 tháng (Từ tháng 11/2026 đến tháng 06/2027)"),
        ("Giảng viên hướng dẫn:", "TS. Nguyễn Hải Hà (Chủ nhiệm Đề tài T2027 TBM_ĐTM)"),
        ("Đơn vị chủ trì:", "Bộ môn Địa kỹ thuật & Không gian ngầm (phối hợp Bộ môn Quản lý Đô thị & Khoa CNTT)"),
        ("Nhóm sinh viên thực hiện:", "1. Sinh viên 1 (Trưởng nhóm - Biên tập Nội dung, Kịch bản StoryMap & Báo cáo)\n2. Sinh viên 2 (Kỹ thuật GIS, Xử lý dữ liệu QGIS, Lập trình WebGIS Low-code/Leaflet)\n3. Sinh viên 3 (Đồ họa UI/UX, Infographic TBM & Tổ chức thử nghiệm Usability)")
    ]
    for r_idx, (k, v) in enumerate(info_data):
        row = tbl.rows[r_idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.5)
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(k)
        format_run(r0, size_pt=11, bold=True)
        set_cell_background(row.cells[0], "F2F2F2")
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(v)
        format_run(r1, size_pt=11)

    add_heading_1(doc, "PHẦN I: TÍNH CẤP THIẾT VÀ TỔNG QUAN VẤN ĐỀ NGHIÊN CỨU")
    add_heading_2(doc, "1.1. Tính cấp thiết của đề tài")
    add_body_p(doc, "Dự án Tuyến đường sắt đô thị số 3 Hà Nội (đoạn ngầm S9 Kim Mã – S12 Ga Hà Nội) dài 4,0 km đang bước vào giai đoạn then chốt với việc vận hành 2 máy đào hầm cân bằng áp lực đất (EPB-TBM) mang tên 'Thần Tốc' và 'Táo Bạo'. Quá trình thi công hầm ngầm đi qua khu vực nội đô lịch sử đông đúc tất yếu làm xuất hiện vùng mất thể tích đất và phễu lún bề mặt (Surface Settlement Trough).", "1. Bối cảnh thực tiễn: ")
    add_body_p(doc, "Dù các báo cáo kỹ thuật địa kỹ thuật và báo cáo ĐTM chứa đựng nhiều dữ liệu chính xác, chúng thường được trình bày dưới dạng công thức giải tích phức tạp (Peck 1969, Loganathan & Poulos 1998, ma trận phần tử hữu hạn). Người dân thông thường không thể tự đọc hiểu và luôn mang tâm lý lo âu: Nhà mình có nằm trong vùng nguy cơ? Cách tim hầm bao nhiêu mét? Khi có vết nứt thì báo cho ai?", "2. Khoảng trống thông tin & nhu cầu xã hội: ")
    add_body_p(doc, "Mô hình GIS có sự tham gia của cộng đồng (Public Participatory GIS - PPGIS) và Bản đồ kể chuyện đa phương tiện (StoryMap) là xu thế quốc tế hàng đầu giúp công khai minh bạch thông tin dự án, giải thích cơ chế kỹ thuật bằng ngôn ngữ trực quan và tạo kênh đối thoại hai chiều giữa Ban quản lý dự án và người dân.", "3. Xu thế Chuyển đổi số & PPGIS: ")
    add_body_p(doc, "Với các công nghệ No-code/Low-code hiện đại (Google My Maps, ArcGIS StoryMaps, Leaflet JS), sinh viên đại học hoàn toàn có khả năng tự tay xây dựng một ứng dụng bản đồ số tương tác chuyên nghiệp, chạy mượt mà trên điện thoại thông minh mà không cần máy chủ đắt tiền hay lập trình phức tạp.", "4. Tính khả thi của giải pháp No-code/Low-code: ")

    add_heading_2(doc, "1.2. Tổng quan tình hình nghiên cứu & Kế thừa Đề tài lớn T2027")
    add_body_p(doc, "Đề tài kế thừa trực tiếp kết quả phân tích giải tích phễu lún Peck (1969) từ Đề tài lớn T2027 của Giảng viên hướng dẫn, đã xác lập 3 dải ranh giới lún quan trọng dọc tuyến S9–S12: (1) Dải 15m - Vùng giám sát đặc biệt (độ lún S > 10–15mm); (2) Dải 35m - Vùng ảnh hưởng phễu lún chính; (3) Dải 50m - Vùng đệm mở rộng an toàn (S < 2–3mm). Đề tài này đưa các dải dữ liệu khoa học đó lên nền tảng bản đồ số tương tác phục vụ cộng đồng.", "Mối liên hệ kế thừa khoa học: ")

    add_heading_1(doc, "PHẦN II: MỤC TIÊU, ĐỐI TƯỢNG VÀ PHẠM VI NGHIÊN CỨU")
    add_body_p(doc, "Xây dựng hệ thống Bản đồ số tương tác (WebGIS) kết hợp Bản đồ kể chuyện (StoryMap) trên nền tảng No-code/Low-code, cung cấp thông tin trực quan về hướng tuyến, tiến độ khiên đào TBM và các vùng dự báo ảnh hưởng lún; đồng thời tích hợp kênh thu nhận phản ánh cộng đồng dọc Tuyến Metro số 3 Hà Nội (đoạn S9–S12).", "Mục tiêu tổng quát: ")
    add_body_p(doc, "1. Thu thập, chuẩn hóa tọa độ không gian WGS-84 và xây dựng bộ cơ sở dữ liệu GeoJSON hoàn chỉnh (tim hầm, 4 ga ngầm, 3 dải đệm lún 15m/35m/50m, các điểm di tích nhạy cảm).\n2. Thiết kế và phát triển Bản đồ số WebGIS tích hợp công cụ tra cứu vị trí nhà dân ('Nhà tôi có an toàn không?'), đo khoảng cách đến tim hầm và hiển thị khuyến nghị an toàn.\n3. Xây dựng Bản đồ kể chuyện StoryMap ('Hành trình máy TBM đi ngầm dưới lòng Hà Nội') bằng đồ họa trực quan và ngôn ngữ đại chúng.\n4. Tích hợp mô-đun tham vấn cộng đồng (PPGIS Feedback) và tổ chức thử nghiệm đánh giá mức độ tiện dụng (thang đo SUS) với 30–50 người dùng.", "Mục tiêu cụ thể: ")

    add_body_p(doc, "Dữ liệu không gian hướng tuyến hầm TBM, ga ngầm, vùng đệm phễu lún và giao diện bản đồ tương tác WebGIS/StoryMap; mức độ tiếp nhận thông tin và phản hồi của người dân.", "Đối tượng nghiên cứu: ")
    add_body_p(doc, "Hành lang đoạn tuyến ngầm Metro số 3 (từ ga S9 Kim Mã qua ga S10 Cát Linh, S11 Văn Miếu đến S12 Ga Hà Nội), mở rộng dải đệm 50m sang hai bên tim hầm.", "Phạm vi không gian: ")

    add_heading_1(doc, "PHẦN III: PHƯƠNG PHÁP NGHIÊN CỨU")
    add_body_p(doc, "1. Phương pháp kế thừa dữ liệu khoa học từ Đề tài T2027 (tọa độ tim hầm và bán kính phễu lún Peck).\n2. Phương pháp mô hình hóa dữ liệu không gian bằng phần mềm mã nguồn mở QGIS (tạo Buffer 15m/35m/50m, xuất định dạng chuẩn GeoJSON/WGS-84).\n3. Phương pháp phát triển WebGIS Low-code/No-code (kết hợp Leaflet JS, Bootstrap 5 và Turf.js cho giải pháp tự chủ; Google My Maps & ArcGIS StoryMaps cho giải pháp no-code nhanh).\n4. Phương pháp truyền thông khoa học bằng câu chuyện (Storytelling 4 màn kịch bản).\n5. Phương pháp đánh giá trải nghiệm người dùng thực tế theo thang đo quốc tế SUS (System Usability Scale).")

    add_heading_1(doc, "PHẦN IV: NỘI DUNG VÀ KẾT CẤU BÁO CÁO (4 CHƯƠNG)")
    add_body_p(doc, "• CHƯƠNG 1: Cơ sở lý luận về GIS cộng đồng (PPGIS) và truyền thông rủi ro trong các dự án hầm đô thị.\n• CHƯƠNG 2: Thu thập, chuẩn hóa và xây dựng cơ sở dữ liệu không gian hành lang tuyến S9–S12.\n• CHƯƠNG 3: Thiết kế và phát triển hệ thống Bản đồ số tương tác (WebGIS) và StoryMap đa phương tiện.\n• CHƯƠNG 4: Thử nghiệm ứng dụng, đánh giá trải nghiệm người dùng và đề xuất giải pháp chuyển đổi số.")

    add_heading_1(doc, "PHẦN V: SẢN PHẨM DỰ KIẾN VÀ ĐÓNG GÓP")
    add_body_p(doc, "1. 01 Báo cáo khoa học toàn văn NCKHSV (50–70 trang) chuẩn mực theo quy định của Trường ĐH GTVT.\n2. 01 Hệ thống Bản đồ số WebGIS tương tác hoàn chỉnh, xuất bản online miễn phí qua GitHub Pages / Netlify.\n3. 01 Kênh StoryMap đa phương tiện 'Hành trình TBM ngầm Thủ đô'.\n4. 01 Bộ cơ sở dữ liệu không gian chuẩn GeoJSON/KML của đoạn hầm ngầm S9–S12.\n5. 01 Bài báo khoa học đăng trên Kỷ yếu Hội nghị Sinh viên Nghiên cứu Khoa học cấp Trường/Khoa năm học 2026–2027.\n6. Đào tạo nhóm sinh viên làm chủ kỹ năng GIS số, tư duy thiết kế sản phẩm và kỹ năng truyền thông khoa học.")

    output_path = os.path.join(TARGET_DIR, "01_DE_CUONG_THUYET_MINH_NCKHSV.docx")
    doc.save(output_path)
    print("Saved:", output_path)

# =========================================================================
# 2. TẠO 02_KE_HOACH_TRIEN_KHAI_CHI_TIET_VA_PHAN_CONG.docx
# =========================================================================
def build_plan_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.79)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("KẾ HOẠCH TRIỂN KHAI CHI TIẾT VÀ PHÂN CÔNG NHIỆM VỤ NCKHSV\n")
    format_run(r_t, size_pt=14, bold=True, color_rgb=(0, 51, 153))
    r_sub = p_title.add_run("ĐỀ TÀI: XÂY DỰNG BẢN ĐỒ SỐ TƯƠNG TÁC (WEBGIS/STORYMAP) VÙNG ẢNH HƯỞNG HẦM METRO HÀ NỘI\nMÃ SỐ: SV-2026-GIS-04 | THỜI GIAN: 08 THÁNG (11/2026 – 06/2027)\n")
    format_run(r_sub, size_pt=11.5, bold=True, color_rgb=(180, 0, 0))

    add_heading_1(doc, "I. MÔ HÌNH VẬN HÀNH: TỔNG THẦU KHOA HỌC & NHÓM KỸ SƯ TRẺ")
    add_body_p(doc, "Định hướng chiến lược, cung cấp dữ liệu gốc từ Đề tài lớn T2027, chủ trì 4 Cổng kiểm soát chất lượng (Quality Gates), hướng dẫn phương pháp học thuật và bảo vệ phản biện.", "Giảng viên hướng dẫn (Tổng thầu khoa học): ")
    add_body_p(doc, "Chịu trách nhiệm tiến độ chung, biên kịch cốt truyện StoryMap, thiết kế câu hỏi tham vấn PPGIS và thang đo SUS, hoàn thiện Báo cáo toàn văn.", "Sinh viên 1 (Trưởng nhóm - Nội dung & PPGIS): ")
    add_body_p(doc, "Chịu trách nhiệm cơ sở dữ liệu không gian QGIS, chuyển hệ tọa độ WGS-84, tạo dải đệm Buffer 15m/35m/50m, lập trình WebGIS Leaflet và thuật toán tra cứu nhà dân.", "Sinh viên 2 (Kỹ thuật GIS & WebGIS): ")
    add_body_p(doc, "Thiết kế giao diện UI/UX tối ưu cho thiết bị di động, thiết kế Infographic 3D máy TBM, điều phối khảo sát thực địa Usability Test với người dân và làm slide/clip bảo vệ.", "Sinh viên 3 (Đồ họa UI/UX & Thử nghiệm): ")

    add_heading_1(doc, "II. MA TRẬN PHÂN CÔNG TRÁCH NHIỆM RACI (12 GÓI CÔNG VIỆC)")
    tbl_raci = doc.add_table(rows=13, cols=6)
    tbl_raci.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_raci)
    raci_headers = ["STT", "Gói công việc (WBS)", "SV1 (Lead)", "SV2 (GIS)", "SV3 (UI)", "GVHD (Duyệt)"]
    for idx, h in enumerate(raci_headers):
        cell = tbl_raci.rows[0].cells[idx]
        cell.paragraphs[0].add_run(h)
        format_run(cell.paragraphs[0].runs[0], size_pt=10, bold=True)
        set_cell_background(cell, "1F4E79")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    raci_rows = [
        ("1", "Hoàn thiện Thuyết minh đề cương chi tiết", "R / A", "C", "C", "Phê duyệt Gate 1"),
        ("2", "Tiếp nhận & kiểm định dữ liệu tọa độ T2027", "C", "R / A", "I", "Bàn giao & Hướng dẫn"),
        ("3", "Xử lý QGIS: Tạo Buffer lún & xuất GeoJSON", "I", "R / A", "C", "Kiểm tra kỹ thuật"),
        ("4", "Biên kịch cốt truyện StoryMap 4 màn", "R / A", "C", "C", "Góp ý nội dung"),
        ("5", "Thiết kế Infographic TBM & Giao diện UI/UX", "C", "C", "R / A", "Góp ý thẩm mỹ"),
        ("6", "Lập trình WebGIS Leaflet & Google My Maps", "C", "R / A", "C", "Đánh giá tính năng"),
        ("7", "Thiết lập form tham vấn cộng đồng PPGIS", "R / A", "C", "C", "Duyệt câu hỏi"),
        ("8", "Tổ chức thử nghiệm Usability Test (30-50 người)", "C", "C", "R / A", "Hỗ trợ điều phối"),
        ("9", "Phân tích điểm SUS & xuất bản Web online", "C", "R", "R", "Nghiệm thu Gate 3"),
        ("10", "Soạn thảo Báo cáo toàn văn (4 Chương)", "R / A", "R", "R", "Phản biện & Sửa chữa"),
        ("11", "Viết bài báo khoa học đăng Kỷ yếu SV", "R / A", "R", "C", "Duyệt & Đồng tác giả"),
        ("12", "Thiết kế Slide, quay Clip demo & Bảo vệ", "C", "C", "R", "Chủ trì duyệt Gate 4")
    ]
    for r_idx, row_vals in enumerate(raci_rows, 1):
        row = tbl_raci.rows[r_idx]
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            is_bold = "R" in val or "Gate" in val or "Duyệt" in val
            format_run(r, size_pt=9.5, bold=is_bold)

    add_heading_1(doc, "III. HỆ THỐNG 4 CỔNG KIỂM SOÁT CHẤT LƯỢNG (AUDIT GATES)")
    add_body_p(doc, "Thuyết minh chi tiết hoàn thiện; file dữ liệu không gian GeoJSON khớp chính xác với hướng tuyến trên bản đồ vệ tinh; kịch bản StoryMap giải thích đúng bản chất khoa học hiện tượng lún.", "Gate 1 (Tuần 5 - Đầu tháng 12/2026): ")
    add_body_p(doc, "Bản mẫu WebGIS hoạt động trơn tru trên máy tính và điện thoại; tính năng tra cứu nhà và đo khoảng cách chính xác; form tham vấn cộng đồng sẵn sàng ghi nhận ý kiến.", "Gate 2 (Tuần 13 - Cuối tháng 01/2027): ")
    add_body_p(doc, "Thu thập tối thiểu 30–50 phiếu thử nghiệm; điểm số tiện dụng SUS > 70 điểm; ứng dụng WebGIS được xuất bản trực tuyến công khai có mã QR quét nhanh.", "Gate 3 (Tuần 21 - Cuối tháng 03/2027): ")
    add_body_p(doc, "Báo cáo toàn văn 50–70 trang không lỗi chính tả; 01 bài báo Kỷ yếu khoa học hoàn chỉnh; slide thuyết trình và video demo 3 phút xuất sắc, nhóm diễn tập tự tin trước hội đồng.", "Gate 4 (Tuần 30 - Cuối tháng 05/2027): ")

    output_path = os.path.join(TARGET_DIR, "02_KE_HOACH_TRIEN_KHAI_CHI_TIET_VA_PHAN_CONG.docx")
    doc.save(output_path)
    print("Saved:", output_path)

# =========================================================================
# 3. TẠO 03_SO_TAY_HUONG_DAN_KY_THUAT_CHO_SINH_VIEN.docx
# =========================================================================
def build_handbook_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.79)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("SỔ TAY HƯỚNG DẪN KỸ THUẬT VÀ CÔNG NGHỆ BẢN ĐỒ SỐ\n")
    format_run(r_t, size_pt=14, bold=True, color_rgb=(0, 51, 153))
    r_sub = p_title.add_run("CẨM NANG \"CẦM TAY CHỈ VIỆC\" CHO SINH VIÊN NCKH WEBGIS & STORYMAP\n(DỰ ÁN BẢN ĐỒ AN TOÀN HẦM METRO LINE 3 HÀ NỘI)\n")
    format_run(r_sub, size_pt=11.5, bold=True, color_rgb=(180, 0, 0))

    add_heading_1(doc, "PHẦN 1: HƯỚNG DẪN XỬ LÝ DỮ LIỆU BẰNG QGIS (CHUẨN HÓA GEOJSON)")
    add_body_p(doc, "1. Cài đặt QGIS miễn phí từ qgis.org. Đặt hệ tọa độ dự án là EPSG:4326 (WGS 84).\n2. Nạp tọa độ các ga ngầm S9..S12 từ bảng Excel bằng chức năng Add Delimited Text Layer.\n3. Dùng công cụ Points to Path để tạo đường tim tuyến hầm ngầm.\n4. Dùng công cụ Vector -> Geoprocessing Tools -> Buffer để tạo 3 dải đệm lún: 15m (vùng giám sát đặc biệt), 35m (vùng phễu lún Peck), 50m (vùng đệm an toàn).\n5. Xuất từng lớp ra định dạng chuẩn GeoJSON (WGS 84) để nhúng trực tiếp vào Web.")

    add_heading_1(doc, "PHẦN 2: TRIỂN KHAI PHƯƠNG ÁN NO-CODE (GOOGLE MY MAPS & ARCGIS STORYMAPS)")
    add_body_p(doc, "Truy cập google.com/mymaps -> Tạo bản đồ mới -> Bấm 'Nhập' file KML -> Thiết lập màu sắc: Đỏ cam (15m), Vàng (35m), Xanh lục (50m) -> Nhúng ảnh thực tế và bật chế độ chia sẻ công khai.", "Phương án 1 - Google My Maps (30 phút hoàn thành): ")
    add_body_p(doc, "Đăng ký tài khoản miễn phí tại storymaps.arcgis.com -> Tạo story mới với mẫu Sidecar -> Tải file GeoJSON lên nền bản đồ vệ tinh ESRI -> Tích hợp hiệu ứng cuộn trang (Scrollytelling) thuyết minh hành trình thi công máy TBM.", "Phương án 2 - ArcGIS Online StoryMaps (Chuẩn quốc tế): ")

    add_heading_1(doc, "PHẦN 3: TRIỂN KHAI PHƯƠNG ÁN LOW-CODE (WEBGIS LEAFLET & BOOTSTRAP)")
    add_body_p(doc, "Đây là giải pháp cốt lõi giúp sinh viên làm chủ mã nguồn độc lập. Sử dụng thư viện Leaflet JS kết hợp Bootstrap 5. Tích hợp thuật toán đo khoảng cách vuông góc từ điểm nhà dân đến tim hầm bằng thư viện Turf.js (turf.pointToLineDistance).")
    add_body_p(doc, "• Nếu khoảng cách <= 15m: Cảnh báo đỏ (Vùng giám sát đặc biệt - dự án gắn mốc quan trắc 24/7).\n• Nếu khoảng cách 15m - 35m: Cảnh báo vàng (Vùng phễu lún Peck - độ lún nhỏ, an toàn kết cấu).\n• Nếu khoảng cách 35m - 50m: Thông báo xanh lục (Vùng đệm an toàn).\n• Nếu khoảng cách > 50m: Thông báo xanh an tâm tuyệt đối (Nằm ngoài vùng ảnh hưởng thi công).")

    add_heading_1(doc, "PHẦN 4: BIÊN KỊCH CỐT TRUYỆN STORYMAP VÀ THAM VẤN PPGIS")
    add_body_p(doc, "Kịch bản 4 màn được biên kịch gần gũi: Màn 1 - Vì sao Thủ đô cần Metro ngầm? Màn 2 - Cận cảnh hai chú robot khổng lồ TBM Thần Tốc và Táo Bạo đào đất lắp hầm như thế nào? Màn 3 - Vì sao mặt đất lại hơi lún nhẹ và dự án kiểm soát an toàn ra sao? Màn 4 - Nhà bạn ở đâu trên hành trình này? Hướng dẫn tự kiểm tra và đường dây nóng hỗ trợ.")
    add_body_p(doc, "Mô-đun tham vấn PPGIS gồm 5 câu hỏi nhanh kết hợp tải ảnh hiện trạng nứt lún; đánh giá khả dụng hệ thống bằng Thang đo tiêu chuẩn quốc tế SUS (10 câu hỏi) hướng tới điểm số > 70 điểm.")

    add_heading_1(doc, "PHẦN 5: XUẤT BẢN TRỰC TUYẾN MIỄN PHÍ VÀ QUẢNG BÁ SẢN PHẨM")
    add_body_p(doc, "Đưa toàn bộ mã nguồn lên GitHub -> Kích hoạt GitHub Pages (miễn phí 100%) -> Nhận đường link web công khai -> Tạo mã QR bằng qr-code-generator.com để in áp phích hoặc chia sẻ vào Zalo tổ dân phố thu thập ý kiến cộng đồng.")

    output_path = os.path.join(TARGET_DIR, "03_SO_TAY_HUONG_DAN_KY_THUAT_CHO_SINH_VIEN.docx")
    doc.save(output_path)
    print("Saved:", output_path)

if __name__ == "__main__":
    build_proposal_docx()
    build_plan_docx()
    build_handbook_docx()
    print("All docx files generated successfully!")
