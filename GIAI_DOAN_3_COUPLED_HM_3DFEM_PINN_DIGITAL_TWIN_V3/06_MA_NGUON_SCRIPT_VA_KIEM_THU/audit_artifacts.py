from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit_workbooks() -> dict:
    error_tokens = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
    records = []
    for path in sorted(ROOT.rglob("*.xlsx")):
        rel = path.relative_to(ROOT).as_posix()
        record = {
            "path": rel,
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
            "status": "READABLE",
        }
        try:
            wb_formula = load_workbook(path, read_only=True, data_only=False, keep_links=True)
            wb_values = load_workbook(path, read_only=True, data_only=True, keep_links=True)
            sheets = []
            formula_count = 0
            formula_cells_without_cached_value = 0
            formula_errors = []
            cached_error_values = []
            for ws_formula, ws_values in zip(wb_formula.worksheets, wb_values.worksheets):
                sheet_formula_count = 0
                for row_formula, row_value in zip(ws_formula.iter_rows(), ws_values.iter_rows()):
                    for cell_formula, cell_value in zip(row_formula, row_value):
                        value = cell_formula.value
                        if isinstance(value, str) and value.startswith("="):
                            formula_count += 1
                            sheet_formula_count += 1
                            if cell_value.value is None:
                                formula_cells_without_cached_value += 1
                            if any(token in value for token in error_tokens):
                                formula_errors.append(f"{ws_formula.title}!{cell_formula.coordinate}:{value}")
                        cached = cell_value.value
                        if isinstance(cached, str) and cached in error_tokens:
                            cached_error_values.append(f"{ws_formula.title}!{cell_value.coordinate}:{cached}")
                sheets.append(
                    {
                        "name": ws_formula.title,
                        "max_row": ws_formula.max_row,
                        "max_column": ws_formula.max_column,
                        "formula_count": sheet_formula_count,
                        "state": ws_formula.sheet_state,
                    }
                )
            calc = wb_formula.calculation
            record.update(
                {
                    "sheets": sheets,
                    "formula_count": formula_count,
                    "formula_cells_without_cached_value": formula_cells_without_cached_value,
                    "formula_text_errors": formula_errors,
                    "cached_error_values": cached_error_values,
                    "calculation": {
                        "calc_mode": getattr(calc, "calcMode", None),
                        "full_calc_on_load": getattr(calc, "fullCalcOnLoad", None),
                        "force_full_calc": getattr(calc, "forceFullCalc", None),
                    },
                    "external_links": len(getattr(wb_formula, "_external_links", [])),
                }
            )
            wb_formula.close()
            wb_values.close()
        except Exception as exc:
            record["status"] = "UNREADABLE"
            record["error"] = f"{type(exc).__name__}: {exc}"
        records.append(record)
    return {
        "workbook_count": len(records),
        "readable_count": sum(r["status"] == "READABLE" for r in records),
        "records": records,
    }


def audit_geojson() -> dict:
    records = []
    for path in sorted(ROOT.rglob("*.geojson")):
        rel = path.relative_to(ROOT).as_posix()
        record = {"path": rel, "status": "READABLE", "size_bytes": path.stat().st_size}
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            features = data.get("features", []) if isinstance(data, dict) else []
            geom_types = Counter()
            qa_flags = Counter()
            source_ids = Counter()
            for feature in features:
                geom = feature.get("geometry") or {}
                geom_types[str(geom.get("type"))] += 1
                props = feature.get("properties") or {}
                qa_flags[str(props.get("qa_flag"))] += 1
                source_ids[str(props.get("source_id"))] += 1
            record.update(
                {
                    "feature_count": len(features),
                    "geometry_types": dict(geom_types),
                    "qa_flags": dict(qa_flags),
                    "source_ids": dict(source_ids),
                    "declared_crs": data.get("crs") if isinstance(data, dict) else None,
                }
            )
        except Exception as exc:
            record["status"] = "UNREADABLE"
            record["error"] = f"{type(exc).__name__}: {exc}"
        records.append(record)
    return {"geojson_count": len(records), "records": records}


def audit_markdown_manifests() -> dict:
    specs = [
        (
            ROOT / "nhiem_vu_luan_an",
            ROOT / "nhiem_vu_luan_an/00_HO_SO_NGHIEM_THU_BO_MON/DANH_MUC_HO_SO_VA_CHECKSUM_SHA256.md",
        ),
        (
            ROOT / "nhiem_vu_luan_an_giai_doan_2",
            ROOT / "nhiem_vu_luan_an_giai_doan_2/06_HO_SO_NGHIEM_THU_GIAI_DOAN_2/DANH_MUC_HO_SO_VA_CHECKSUM_SHA256_GIAI_DOAN_2.md",
        ),
        (
            ROOT / "nhiem_vu_luan_an_giai_doan_3",
            ROOT / "nhiem_vu_luan_an_giai_doan_3/06_HO_SO_NGHIEM_THU_GIAI_DOAN_3/DANH_MUC_HO_SO_VA_CHECKSUM_SHA256_GIAI_DOAN_3.md",
        ),
    ]
    pattern = re.compile(r"\|\s*\d+\s*\|\s*`([^`]+)`\s*\|\s*([^|]+)\|\s*`([0-9a-fA-F]{64})`")
    reports = []
    for root, manifest in specs:
        entries = []
        for match in pattern.finditer(manifest.read_text(encoding="utf-8")):
            rel, size_text, expected = match.groups()
            path = root / Path(rel)
            item = {"path": rel, "expected_sha256": expected.lower(), "declared_size": size_text.strip()}
            if not path.exists():
                item["status"] = "MISSING"
            else:
                try:
                    actual = sha256(path)
                    item["actual_sha256"] = actual
                    item["actual_size_bytes"] = path.stat().st_size
                    item["status"] = "MATCH" if actual == expected.lower() else "MISMATCH"
                except PermissionError as exc:
                    item["status"] = "LOCKED_UNVERIFIED"
                    item["error"] = str(exc)
            entries.append(item)
        reports.append(
            {
                "manifest": manifest.relative_to(ROOT).as_posix(),
                "entry_count": len(entries),
                "status_counts": dict(Counter(item["status"] for item in entries)),
                "issues": [item for item in entries if item["status"] != "MATCH"],
            }
        )
    return {"reports": reports}


def main() -> None:
    result = {
        "audit_scope": "read-only artifact structure, workbook internals, GeoJSON metadata, and declared SHA-256 manifests",
        "project_root": str(ROOT),
        "workbooks": audit_workbooks(),
        "geojson": audit_geojson(),
        "manifests": audit_markdown_manifests(),
    }
    output = OUT / "artifact_audit.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(output.relative_to(ROOT).as_posix())


if __name__ == "__main__":
    main()
